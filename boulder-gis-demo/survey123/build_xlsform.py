"""Build the Survey123 XLSForm for the Boulder Permit Site Check.
Open the output in Survey123 Connect, then publish. Requires openpyxl."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font

SURVEY = [
    # type, name, label, hint, required, relevant, constraint, constraint_message, default, appearance, calculation
    ("begin group", "permit_info", "Permit", "", "", "", "", "", "", "field-list", ""),
    ("text", "permit_num", "Permit number", "Filled in from the Field Maps pop-up", "yes", "",
     "regex(., '^[A-Z]{3}\\d{4}-\\d{5}$')", "Use the format PMT2026-00001", "", "", ""),
    ("text", "address", "Site address", "", "", "", "", "", "", "", ""),
    ("geopoint", "site_location", "Site location", "Confirm or move the point", "yes", "", "", "", "", "", ""),
    ("end group", "", "", "", "", "", "", "", "", "", ""),
    ("begin group", "visit", "Site visit", "", "", "", "", "", "", "field-list", ""),
    ("date", "visit_date", "Visit date", "", "yes", "", ". <= today()", "Visit date cannot be in the future",
     "today", "", ""),
    ("username", "inspector", "Inspector", "", "", "", "", "", "", "", ""),
    ("select_one site_status", "site_status", "Site status", "", "yes", "", "", "", "", "", ""),
    ("select_one yes_no_na", "permit_card_posted", "Is the permit card posted and visible?", "", "yes",
     "${site_status} = 'active_construction'", "", "", "", "horizontal", ""),
    ("select_one erosion", "erosion_control", "Erosion and sediment control", "", "yes",
     "${site_status} = 'active_construction'", "", "", "", "", ""),
    ("image", "erosion_photo", "Photo of erosion control issue", "", "yes",
     "${erosion_control} = 'deficient'", "", "", "", "", ""),
    ("select_one yes_no", "row_obstruction", "Is the sidewalk or right-of-way blocked?", "", "yes",
     "${site_status} = 'active_construction'", "", "", "", "horizontal", ""),
    ("image", "row_photo", "Photo of obstruction", "", "yes", "${row_obstruction} = 'yes'", "", "", "", "", ""),
    ("text", "notes", "Notes", "", "", "", "", "", "", "multiline", ""),
    ("end group", "", "", "", "", "", "", "", "", "", ""),
    ("calculate", "followup_needed", "Follow-up needed", "", "", "", "", "", "", "",
     "if(${permit_card_posted} = 'no' or ${erosion_control} = 'deficient' or ${row_obstruction} = 'yes', 'yes', 'no')"),
]

CHOICES = [
    ("site_status", "active_construction", "Active construction"),
    ("site_status", "not_started", "Not started"),
    ("site_status", "complete", "Work appears complete"),
    ("site_status", "no_access", "No access"),
    ("yes_no_na", "yes", "Yes"), ("yes_no_na", "no", "No"), ("yes_no_na", "na", "N/A"),
    ("yes_no", "yes", "Yes"), ("yes_no", "no", "No"),
    ("erosion", "in_place", "In place"),
    ("erosion", "deficient", "Deficient"),
    ("erosion", "not_required", "Not required"),
]

wb = Workbook()
ws = wb.active
ws.title = "survey"
ws.append(["type", "name", "label", "hint", "required", "relevant", "constraint", "constraint_message",
           "default", "appearance", "calculation"])
for row in SURVEY:
    ws.append(list(row))

wc = wb.create_sheet("choices")
wc.append(["list_name", "name", "label"])
for row in CHOICES:
    wc.append(list(row))

wset = wb.create_sheet("settings")
wset.append(["form_title", "form_id", "instance_name", "style"])
wset.append(["Boulder Permit Site Check", "boulder_permit_site_check",
             "concat(${permit_num}, ' - ', ${visit_date})", "pages"])

for sheet in wb.worksheets:
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 24

out = Path(__file__).parent / "boulder_permit_site_check.xlsx"
wb.save(out)
print(f"wrote {out.name}")
